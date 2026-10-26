"""Bounded, deterministic CSV/XLSX/JSON/Parquet schema extraction."""

from __future__ import annotations

import csv
import io
import json
import math
import zipfile
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from etl_agent.modules.assets.ports import ParsedFile
from etl_agent.shared.errors import ApplicationError, invalid

MAX_BYTES = 100 * 1024 * 1024


def field(
    name: str,
    ordinal: int,
    kind: str,
    *,
    nullable: bool | None = None,
    native_type: str | None = None,
    precision: int | None = None,
    scale: int | None = None,
) -> dict[str, Any]:
    return {
        "field_name": name,
        "ordinal": ordinal,
        "native_type": native_type,
        "logical_type": kind,
        "nullable": nullable,
        "is_primary_key": False,
        "precision": precision,
        "scale": scale,
        "sensitivity": "unclassified",
    }


def headers(values: list[Any]) -> list[str]:
    if not values or any(not isinstance(item, str) or not item.strip() for item in values):
        raise invalid("文件必须有明确且非空的文本列名", "file")
    if len(set(values)) != len(values):
        raise invalid("文件包含重复列名", "file")
    return values


def kind(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, (float, Decimal)):
        if not math.isfinite(float(value)):
            raise invalid("不接受非有限数值", "file")
        return "decimal"
    if isinstance(value, datetime):
        return "timestamp"
    if isinstance(value, date):
        return "date"
    if isinstance(value, str):
        return "string"
    if isinstance(value, bytes):
        return "binary"
    if isinstance(value, list):
        return "array"
    if isinstance(value, dict):
        return "object"
    raise invalid("文件包含无法确定类型的值", "file")


def infer(names: list[str], rows: list[list[Any]]) -> ParsedFile:
    result = []
    for ordinal, name in enumerate(names):
        values = [row[ordinal] for row in rows]
        actual = [item for item in values if item is not None]
        kinds = {kind(item) for item in actual}
        if not kinds:
            selected = "unknown"
        elif kinds.issubset({"integer", "decimal"}):
            selected = "decimal" if "decimal" in kinds else "integer"
        elif len(kinds) == 1:
            selected = next(iter(kinds))
        else:
            raise invalid("同列存在歧义类型，需要先规范源文件", "file")
        precision = scale = None
        if selected == "decimal":
            decimals = [Decimal(str(item)) for item in actual]
            scale = max(max(-int(item.as_tuple().exponent), 0) for item in decimals)
            whole = max(max(item.adjusted() + 1, 0) for item in decimals)
            precision = max(whole + scale, 1)
        result.append(
            field(
                name,
                ordinal,
                selected,
                nullable=any(item is None for item in values),
                precision=precision,
                scale=scale,
            )
        )
    return ParsedFile(result, {"row_count": str(len(rows)), "is_approximate": False})


class StructuredFileParser:
    def parse(self, content: bytes, file_format: str) -> ParsedFile:
        if not content:
            raise invalid("空文件没有可解析Schema", "file")
        if len(content) > MAX_BYTES:
            raise invalid("文件超过100MiB", "file")
        try:
            parser = getattr(self, f"_{file_format}", None)
            if parser is None or file_format not in {"csv", "xlsx", "json", "parquet"}:
                raise invalid("不支持该文件格式", "file")
            return parser(content)
        except ApplicationError:
            raise
        except Exception as error:
            raise invalid("文件内容损坏或与声明格式不符", "file") from error

    def _csv(self, content: bytes) -> ParsedFile:
        value = content.decode("utf-8-sig", errors="strict")
        # CSV has text values. Guessing numeric types would corrupt leading zeroes.
        dialect = csv.Sniffer().sniff(value[:8192], delimiters=",;\t|")
        reader = csv.reader(io.StringIO(value), dialect, strict=True)
        names = headers(next(reader))
        count = 0
        for row in reader:
            if len(row) != len(names):
                raise invalid("CSV数据列数与表头不一致", "file")
            count += 1
        return ParsedFile(
            [field(name, index, "string", nullable=False) for index, name in enumerate(names)],
            {"row_count": str(count), "is_approximate": False},
        )

    def _json(self, content: bytes) -> ParsedFile:
        def unique_object(pairs):
            output = {}
            for key, value in pairs:
                if key in output:
                    raise invalid("JSON对象包含重复字段", "file")
                output[key] = value
            return output

        def reject_constant(value):
            raise invalid("JSON包含非有限数值", "file")

        records = json.loads(
            content.decode("utf-8-sig"),
            parse_float=Decimal,
            parse_constant=reject_constant,
            object_pairs_hook=unique_object,
        )
        if (
            not isinstance(records, list)
            or not records
            or not all(isinstance(item, dict) for item in records)
        ):
            raise invalid("JSON须为非空对象数组，不能猜测根对象数据路径", "file")
        names = headers(list(records[0]))
        if any(set(item) != set(names) for item in records):
            raise invalid("JSON记录字段不一致，需要先规范源文件", "file")
        return infer(names, [[record[name] for name in names] for record in records])

    def _xlsx(self, content: bytes) -> ParsedFile:
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_BYTES:
                raise invalid("Excel解压内容超过100MiB", "file")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.sheetnames) != 1:
                raise invalid("Excel必须只有一个工作表，避免猜测数据范围", "file")
            rows = workbook.active.iter_rows()
            names = headers([cell.value for cell in next(rows)])
            values = []
            for row in rows:
                if len(row) != len(names):
                    raise invalid("Excel数据列数不一致", "file")
                if any(cell.data_type == "f" for cell in row):
                    raise invalid("Excel包含未计算公式，需要先导出为确定值", "file")
                values.append([cell.value for cell in row])
            return infer(names, values)
        finally:
            workbook.close()

    def _parquet(self, content: bytes) -> ParsedFile:
        import pyarrow as pa
        import pyarrow.parquet as pq

        source = pq.ParquetFile(io.BytesIO(content))
        schema = source.schema_arrow
        names = headers(schema.names)
        fields = []
        for ordinal, item in enumerate(schema):
            value = item.type
            if pa.types.is_decimal(value):
                logical = "decimal"
            elif pa.types.is_integer(value):
                logical = "integer"
            elif pa.types.is_floating(value):
                # No fake precision/scale for a binary floating type.
                raise invalid("Parquet浮点列缺少精确decimal契约，请先转换源字段", "file")
            elif pa.types.is_boolean(value):
                logical = "boolean"
            elif pa.types.is_timestamp(value):
                logical = "timestamp"
            elif pa.types.is_date(value):
                logical = "date"
            elif pa.types.is_string(value) or pa.types.is_large_string(value):
                logical = "string"
            elif pa.types.is_binary(value) or pa.types.is_large_binary(value):
                logical = "binary"
            elif pa.types.is_list(value) or pa.types.is_large_list(value):
                logical = "array"
            elif pa.types.is_struct(value) or pa.types.is_map(value):
                logical = "object"
            else:
                raise invalid("Parquet包含未支持的字段类型", "file")
            fields.append(
                field(
                    names[ordinal],
                    ordinal,
                    logical,
                    nullable=item.nullable,
                    native_type=str(value),
                    precision=value.precision if logical == "decimal" else None,
                    scale=value.scale if logical == "decimal" else None,
                )
            )
        return ParsedFile(
            fields, {"row_count": str(source.metadata.num_rows), "is_approximate": False}
        )
